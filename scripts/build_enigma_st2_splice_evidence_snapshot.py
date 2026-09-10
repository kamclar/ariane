"""Build the complete ENIGMA ST2 splice evidence and ST3 reference snapshot."""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import date
from pathlib import Path

import openpyxl


SOURCE_URL = (
    "https://cspec.genome.network/cspec/File/id/"
    "cb4a09fe-30f4-4aa8-9d76-d7ea407c9754/data"
)
SHEET = "ST2 splicing dataset codes"
REFERENCE_SHEET = "ST3 Splicing refs"
FIELDS = [
    "gene",
    "c_notation",
    "p_notation",
    "agvgd_grade",
    "key_domain",
    "included_in_analysis",
    "splicing_assay_result_category",
    "variant_assay_summary",
    "prior_probability",
    "final_multifactorial_class",
    "result",
]
REFERENCE_FIELDS = [
    "gene",
    "c_notation",
    "p_notation",
    "final_class",
    "author",
    "year",
    "pmid",
    "patient_rna",
    "minigene",
    "result",
    "aberration_type",
    "rna_change",
    "wild_type_transcript_produced",
    "percent_aberrant_transcript_reported",
    "allele_specific",
    "comment",
]


def _clean(value):
    if isinstance(value, str):
        return value.strip()
    return value


def build_snapshot(source: Path) -> dict:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheet = workbook[SHEET]
    reference_sheet = workbook[REFERENCE_SHEET]
    references_by_variant = {}
    reference_count = 0
    for source_row, row in enumerate(
        reference_sheet.iter_rows(min_row=4, values_only=True), 4
    ):
        if row[0] not in {"BRCA1", "BRCA2"}:
            continue
        reference = {
            field: _clean(value)
            for field, value in zip(REFERENCE_FIELDS, row[: len(REFERENCE_FIELDS)])
        }
        reference["source_row"] = source_row
        reference["final_class"] = int(reference["final_class"])
        references_by_variant.setdefault(
            (reference["gene"], reference["c_notation"]), []
        ).append(reference)
        reference_count += 1

    if reference_count != 383:
        raise RuntimeError(f"Expected 383 ST3 reference rows, found {reference_count}")

    variants = []
    seen = set()
    for source_row, row in enumerate(sheet.iter_rows(min_row=6, values_only=True), 6):
        if row[0] not in {"BRCA1", "BRCA2"}:
            continue
        entry = {
            field: _clean(value)
            for field, value in zip(FIELDS, row[: len(FIELDS)])
        }
        entry["source_row"] = source_row
        entry["final_multifactorial_class"] = int(
            entry["final_multifactorial_class"]
        )
        key = (entry["gene"], entry["c_notation"])
        if key in seen:
            raise RuntimeError(f"Duplicate ST2 variant: {key}")
        seen.add(key)
        entry["st3_references"] = references_by_variant.get(key, [])
        if not entry["st3_references"]:
            raise RuntimeError(f"ST2 variant has no supporting ST3 reference: {key}")
        variants.append(entry)

    if len(variants) != 220:
        raise RuntimeError(f"Expected 220 ST2 variants, found {len(variants)}")
    return {
        "schema_version": 2,
        "version": "1.2.0",
        "released": "2025-01-09",
        "generated": date.today().isoformat(),
        "source": "ClinGen ENIGMA BRCA1/2 VCEP Supplementary Tables 2 and 3",
        "source_url": SOURCE_URL,
        "source_file": source.name,
        "source_file_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "source_sheet": SHEET,
        "source_columns": len(FIELDS),
        "reference_source_sheet": REFERENCE_SHEET,
        "reference_source_columns": len(REFERENCE_FIELDS),
        "total_reference_rows": reference_count,
        "total_variants": len(variants),
        "absence_semantics": (
            "No exact row means only that no record was identified in this versioned "
            "ENIGMA ST2 snapshot; it is not proof that no splice evidence exists."
        ),
        "variants": variants,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("backend/data/enigma_st2_splice_evidence.json"),
    )
    args = parser.parse_args()
    snapshot = build_snapshot(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"Wrote {snapshot['total_variants']} complete ST2 rows to {args.output}")


if __name__ == "__main__":
    main()
